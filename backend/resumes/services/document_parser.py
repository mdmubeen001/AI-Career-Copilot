import io
import os
import pypdf
import docx
from django.conf import settings


class DocumentParsingError(ValueError):
    """Custom exception raised when document parsing fails."""
    pass


class DocumentParser:
    """
    Parser service extracting raw text from PDF and DOCX documents safely.
    Validates file formats, sizes, and integrity.
    """

    ALLOWED_EXTENSIONS = {'.pdf': 'pdf', '.docx': 'docx'}

    @classmethod
    def parse_file(cls, uploaded_file) -> tuple:
        """
        Convenience method to validate and extract readable text in one step.
        Returns:
            (extracted_text, file_type)
        """
        file_type = cls.validate_file(uploaded_file)
        extracted_text = cls.extract_text(uploaded_file, file_type)
        return extracted_text, file_type

    @classmethod
    def validate_file(cls, uploaded_file) -> str:
        """
        Validate file extension, non-empty content, and file size.
        Returns normalized file_type ('pdf' or 'docx').
        """
        if not uploaded_file or not hasattr(uploaded_file, 'name'):
            raise DocumentParsingError("No valid file provided.")

        ext = os.path.splitext(uploaded_file.name)[1].lower()
        if ext not in cls.ALLOWED_EXTENSIONS:
            raise DocumentParsingError(
                f"Unsupported file format '{ext}'. Only PDF (.pdf) and Word (.docx) documents are accepted."
            )

        max_size = getattr(settings, 'MAX_UPLOAD_SIZE', 5 * 1024 * 1024)
        if uploaded_file.size > max_size:
            max_mb = max_size // (1024 * 1024)
            raise DocumentParsingError(
                f"File exceeds the maximum allowed size of {max_mb} MB."
            )

        if uploaded_file.size == 0:
            raise DocumentParsingError("Uploaded file is empty.")

        return cls.ALLOWED_EXTENSIONS[ext]

    @classmethod
    def extract_text(cls, uploaded_file, file_type: str) -> str:
        """
        Extract readable plain text from the uploaded document.
        """
        # Ensure pointer is at beginning of stream
        if hasattr(uploaded_file, 'seek'):
            uploaded_file.seek(0)

        file_bytes = uploaded_file.read()
        if hasattr(uploaded_file, 'seek'):
            uploaded_file.seek(0)

        stream = io.BytesIO(file_bytes)

        if file_type == 'pdf':
            text = cls._extract_from_pdf(stream)
        elif file_type == 'docx':
            text = cls._extract_from_docx(stream)
        else:
            raise DocumentParsingError(f"Unsupported file type: {file_type}")

        cleaned_text = text.strip()
        if not cleaned_text:
            raise DocumentParsingError(
                "Could not extract readable text from the document. "
                "Scanned/image-only PDFs or empty documents are not supported. "
                "Please ensure your resume contains selectable text."
            )

        return cleaned_text

    @classmethod
    def _extract_from_pdf(cls, stream: io.BytesIO) -> str:
        try:
            reader = pypdf.PdfReader(stream)
            if reader.is_encrypted:
                try:
                    # Attempt empty password decryption
                    reader.decrypt("")
                except Exception:
                    raise DocumentParsingError("Password-protected or encrypted PDFs are not supported.")

            pages_text = []
            for idx, page in enumerate(reader.pages):
                try:
                    page_content = page.extract_text() or ""
                    if page_content.strip():
                        pages_text.append(page_content)
                except Exception:
                    continue

            return "\n\n".join(pages_text)
        except Exception as e:
            if isinstance(e, DocumentParsingError):
                raise e
            raise DocumentParsingError(f"Corrupted or invalid PDF file: {str(e)}")

    @classmethod
    def _extract_from_docx(cls, stream: io.BytesIO) -> str:
        try:
            doc = docx.Document(stream)
            parts = []

            for p in doc.paragraphs:
                if p.text.strip():
                    parts.append(p.text.strip())

            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        parts.append(" | ".join(row_cells))

            return "\n".join(parts)
        except Exception as e:
            raise DocumentParsingError(f"Corrupted or invalid DOCX document: {str(e)}")
